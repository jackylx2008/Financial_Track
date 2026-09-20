from __future__ import annotations

import json
import csv
import re
from io import StringIO
from pathlib import Path
from typing import Any

from flows.modules.bank_transaction_schema import make_transaction, parse_money_token
from flows.modules.financial_document_ai import FinancialDocumentAiFallback


DATE_LINE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
TIME_RE = re.compile(r"^\d{2}:\d{2}:\d{2}$")
ACCOUNT_RE = re.compile(r"(?:卡号|账号|卡号/账号)[:： ]*([0-9*]{4,})")
CARD_TAIL_RE = re.compile(r"\(([0-9*]{4,})\)")
SIGNED_AMOUNT_RE = re.compile(r"^[+-]\d{1,3}(?:,\d{3})*(?:\.\d{2})$|^[+-]\d+(?:\.\d{2})$")


def read_attachment_transactions(
    manifest_path: Path,
    ai_fallback: FinancialDocumentAiFallback | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if not manifest_path.exists():
        return [], {
            "manifest_exists": False,
            "files_seen": 0,
            "transactions": 0,
            "parse_failures": [],
            "unresolved_files": [],
        }
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    transactions: list[dict[str, Any]] = []
    failures: list[str] = []
    pending_ai: list[tuple[Path, dict[str, Any]]] = []
    unresolved_files: list[dict[str, str]] = []
    files_seen = 0
    for item in manifest:
        if item.get("status") != "success":
            continue
        for output_file in item.get("output_files", []):
            files_seen += 1
            path = Path(output_file)
            try:
                parsed: list[dict[str, Any]] = []
                if path.suffix.lower() == ".pdf":
                    parsed = _read_pdf_transactions(path, item)
                elif path.suffix.lower() == ".xls":
                    parsed = _read_xls_transactions(path, item)
                elif path.suffix.lower() == ".xlsx":
                    parsed = _read_xlsx_transactions(path, item)
                elif path.suffix.lower() == ".csv":
                    parsed = _read_csv_transactions(path, item)
                if not parsed:
                    if ai_fallback is not None and path.suffix.lower() in {".pdf", ".xls", ".xlsx", ".csv"}:
                        pending_ai.append((path, item))
                    else:
                        unresolved_files.append(_unresolved_file(path, item))
                transactions.extend(parsed)
            except Exception as exc:
                failures.append(f"{path}: {exc}")
                unresolved_files.append(_unresolved_file(path, item, f"自动解析失败：{type(exc).__name__}"))
    priority = {".pdf": 0, ".xls": 1, ".xlsx": 1, ".csv": 2}
    for path, item in sorted(pending_ai, key=lambda pair: (priority.get(pair[0].suffix.lower(), 9), str(pair[0]))):
        try:
            ai_rows = _read_attachment_with_ai(path, item, ai_fallback)
            transactions.extend(ai_rows)
            if not ai_rows:
                unresolved_files.append(_unresolved_file(path, item))
        except Exception as exc:
            failures.append(f"{path}: AI fallback: {type(exc).__name__}: {exc}")
            unresolved_files.append(_unresolved_file(path, item, f"AI/OCR 处理失败：{type(exc).__name__}"))
    for transaction in transactions:
        _apply_known_card_metadata(transaction)
    stats = {
        "manifest_exists": True,
        "files_seen": files_seen,
        "transactions": len(transactions),
        "parse_failures": failures,
        "unresolved_files": unresolved_files,
    }
    if ai_fallback is not None:
        stats["ai_fallback"] = ai_fallback.stats()
    return transactions, stats


def _unresolved_file(
    path: Path,
    manifest_item: dict[str, Any],
    reason: str = "自动解析未提取到交易",
) -> dict[str, str]:
    bank_key = str(manifest_item.get("bank_key", "unknown") or "unknown")
    return {
        "bank_key": bank_key,
        "bank_name": str(manifest_item.get("bank_name", "")).strip() or _bank_name(bank_key) or "未知机构",
        "file_type": path.suffix.lower().lstrip(".").upper() or "文件",
        "filename": path.name,
        "path": str(path),
        "reason": reason,
    }


def _read_pdf_transactions(path: Path, manifest_item: dict[str, Any]) -> list[dict[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    text = "\n".join(_fix_mojibake(page.extract_text() or "") for page in reader.pages)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    account_tail = _account_tail_from_lines(lines)
    account_full_name = _account_full_from_lines(lines)
    transactions: list[dict[str, Any]] = []
    for index, line in enumerate(lines[:-1]):
        if not DATE_LINE_RE.match(line):
            continue
        next_line = lines[index + 1]
        if not TIME_RE.match(next_line.split()[0] if next_line.split() else ""):
            continue
        tx = _parse_icbc_pdf_line(
            date=line,
            detail_line=next_line,
            account_tail=account_tail,
            account_full_name=account_full_name,
            path=path,
            manifest_item=manifest_item,
            page_hint="",
        )
        if tx is not None:
            transactions.append(tx)
    return transactions


def _read_csv_transactions(path: Path, manifest_item: dict[str, Any]) -> list[dict[str, Any]]:
    text = _read_csv_text(path)
    rows = _read_dict_rows_from_text(text)
    if not rows:
        return []
    headers = set(rows[0])
    if {"交易时间", "交易分类", "交易对方", "收/支", "金额"}.issubset(headers):
        return _read_alipay_csv_rows(path, manifest_item, rows)
    if {"交易创建时间", "交易成功时间", "订单标题", "收/支", "实付金额"}.issubset(headers):
        return _read_meituan_csv_rows(path, manifest_item, rows)
    return _read_generic_bank_rows(path, manifest_item, rows, "email_attachment_csv")


def _read_alipay_csv_rows(
    path: Path, manifest_item: dict[str, Any], rows: list[dict[str, str]]
) -> list[dict[str, Any]]:
    transactions: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows, start=1):
        amount = _clean_money(row.get("金额", ""))
        if not parse_money_token(amount):
            continue
        direction = _direction_from_cn(row.get("收/支", ""))
        transactions.append(
            make_transaction(
                bank_key="alipay",
                bank_name="支付宝",
                account_full_name=row.get("收/付款方式", "").strip(),
                account_tail=_account_tail_from_text(row.get("收/付款方式", "")),
                transaction_time=row.get("交易时间", "").strip(),
                direction=direction,
                amount=amount,
                merchant=row.get("交易对方", "").strip(),
                counterparty=row.get("交易对方", "").strip(),
                summary=row.get("商品说明", "").strip() or row.get("交易分类", "").strip(),
                channel=row.get("收/付款方式", "").strip(),
                transaction_reference=row.get("交易订单号", "").strip(),
                source_records=[
                    {
                        "source_type": "email_attachment_csv",
                        "source_file": str(path),
                        "message_uid": manifest_item.get("message_uid", ""),
                        "message_id": manifest_item.get("message_id", ""),
                        "row": row_index,
                    }
                ],
                confidence=0.9,
                raw_record=row,
            )
        )
    return transactions


def _read_meituan_csv_rows(
    path: Path, manifest_item: dict[str, Any], rows: list[dict[str, str]]
) -> list[dict[str, Any]]:
    transactions: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows, start=1):
        amount = _clean_money(row.get("实付金额", "") or row.get("订单金额", ""))
        if not parse_money_token(amount):
            continue
        transactions.append(
            make_transaction(
                bank_key="meituan",
                bank_name="美团",
                account_full_name=row.get("支付方式", "").strip(),
                account_tail=_account_tail_from_text(row.get("支付方式", "")),
                transaction_time=(row.get("交易成功时间", "") or row.get("交易创建时间", "")).strip(),
                direction=_direction_from_cn(row.get("收/支", "")),
                amount=amount,
                merchant=row.get("订单标题", "").strip(),
                counterparty=row.get("订单标题", "").strip(),
                summary=row.get("交易类型", "").strip(),
                channel=row.get("支付方式", "").strip(),
                transaction_reference=row.get("交易单号", "").strip(),
                source_records=[
                    {
                        "source_type": "email_attachment_csv",
                        "source_file": str(path),
                        "message_uid": manifest_item.get("message_uid", ""),
                        "message_id": manifest_item.get("message_id", ""),
                        "row": row_index,
                    }
                ],
                confidence=0.9,
                raw_record=row,
            )
        )
    return transactions


def _read_csv_text(path: Path) -> str:
    data = path.read_bytes()
    best_text = ""
    best_score = -1
    for encoding in ("utf-8-sig", "gb18030", "gbk", "utf-16", "latin1"):
        try:
            decoded = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        fixed = "\n".join(_fix_mojibake(line) for line in decoded.splitlines())
        score = _score_cjk(fixed)
        if score > best_score:
            best_text = fixed
            best_score = score
    return best_text


def _read_dict_rows_from_text(text: str) -> list[dict[str, str]]:
    lines = [line for line in text.splitlines() if line.strip().strip('"')]
    for index, line in enumerate(lines):
        if not any(label in line for label in ("交易时间", "交易创建时间", "交易日期", "发生时间", "记账时间")):
            continue
        reader = csv.DictReader(StringIO("\n".join(lines[index:])))
        return [
            {str(key or "").strip(): str(value or "").strip() for key, value in row.items()}
            for row in reader
        ]
    return []


def _parse_icbc_pdf_line(
    date: str,
    detail_line: str,
    account_tail: str,
    account_full_name: str,
    path: Path,
    manifest_item: dict[str, Any],
    page_hint: str,
) -> dict[str, Any] | None:
    tokens = detail_line.split()
    if len(tokens) < 8:
        return None
    if tokens[2] in {"借", "贷"}:
        return _parse_icbc_credit_card_line(
            date,
            tokens,
            account_tail,
            account_full_name,
            path,
            manifest_item,
            page_hint,
            detail_line,
        )
    amount_index = next((idx for idx, token in enumerate(tokens) if SIGNED_AMOUNT_RE.match(token)), -1)
    if amount_index < 0:
        return None
    amount = tokens[amount_index]
    balance = tokens[amount_index + 1] if amount_index + 1 < len(tokens) else ""
    summary_tokens = tokens[6 : max(6, amount_index - 1)]
    trailing_fields = tokens[amount_index + 2 :]
    counterparty = trailing_fields[0] if trailing_fields else ""
    counterparty_account = trailing_fields[1] if len(trailing_fields) >= 2 else ""
    channel = trailing_fields[-1] if len(trailing_fields) >= 3 else tokens[amount_index - 1]
    return make_transaction(
        bank_key=str(manifest_item.get("bank_key", "icbc") or "icbc"),
        bank_name="工商银行",
        account_full_name=account_full_name,
        account_tail=account_tail,
        transaction_time=f"{date} {tokens[0]}",
        direction="inflow" if amount.startswith("+") else "outflow",
        amount=amount,
        summary=" ".join(summary_tokens),
        balance=balance,
        counterparty=counterparty,
        merchant=counterparty,
        counterparty_account=counterparty_account,
        channel=channel,
        transaction_type=tokens[amount_index - 1] if amount_index >= 1 else "",
        source_records=[
            {
                "source_type": "email_attachment_pdf",
                "source_file": str(path),
                "message_uid": manifest_item.get("message_uid", ""),
                "message_id": manifest_item.get("message_id", ""),
                "page": page_hint,
            }
        ],
        confidence=0.82,
        raw_record={"line": detail_line},
    )


def _parse_icbc_credit_card_line(
    date: str,
    tokens: list[str],
    account_tail: str,
    account_full_name: str,
    path: Path,
    manifest_item: dict[str, Any],
    page_hint: str,
    detail_line: str,
) -> dict[str, Any] | None:
    """解析工行信用卡明细，避免把带符号的账户余额当作交易金额。"""
    if len(tokens) < 8:
        return None
    transaction_amount = parse_money_token(tokens[4])
    posting_amount = parse_money_token(tokens[6])
    balance = parse_money_token(tokens[7])
    amount = transaction_amount or posting_amount
    if not amount:
        return None
    direction = "outflow" if tokens[2] == "借" else "inflow"
    summary = tokens[8] if len(tokens) > 8 else ""
    counterparty = " ".join(tokens[9:])
    transaction = make_transaction(
        bank_key=str(manifest_item.get("bank_key", "icbc") or "icbc"),
        bank_name="工商银行",
        account_full_name=account_full_name,
        account_tail=account_tail,
        transaction_time=f"{date} {tokens[0]}",
        posting_date=date,
        direction=direction,
        amount=amount,
        currency=tokens[3],
        merchant=counterparty,
        counterparty=counterparty,
        summary=summary,
        balance=balance,
        transaction_type=tokens[2],
        source_records=[
            {
                "source_type": "email_attachment_pdf",
                "source_file": str(path),
                "message_uid": manifest_item.get("message_uid", ""),
                "message_id": manifest_item.get("message_id", ""),
                "page": page_hint,
            }
        ],
        confidence=0.9,
        raw_record={"line": detail_line},
    )
    transaction["card_type"] = "信用卡"
    transaction["transaction_card_tail"] = re.sub(r"\D", "", tokens[1])[-4:]
    return transaction


def _apply_known_card_metadata(transaction: dict[str, Any]) -> None:
    bank_key = str(transaction.get("bank_key", ""))
    if bank_key in {"bocom", "ccb"}:
        transaction["card_type"] = "借记卡"
    elif bank_key == "cmb":
        transaction["card_type"] = "信用卡"
    elif bank_key == "icbc" and not transaction.get("card_type"):
        raw_record = transaction.get("raw_record")
        raw_line = str(raw_record.get("line", "")) if isinstance(raw_record, dict) else ""
        tokens = raw_line.split()
        transaction["card_type"] = "信用卡" if len(tokens) >= 3 and tokens[2] in {"借", "贷"} else "借记卡"


def _read_xls_transactions(path: Path, manifest_item: dict[str, Any]) -> list[dict[str, Any]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    transactions: list[dict[str, Any]] = []
    for sheet in book.sheets():
        header_row = _find_header_row(sheet)
        if header_row is None:
            continue
        headers = [str(sheet.cell_value(header_row, col)).strip() for col in range(sheet.ncols)]
        account_full_name = _account_full_from_xls(sheet)
        account_tail = _account_tail_from_text(account_full_name)
        for row_index in range(header_row + 1, sheet.nrows):
            row = {headers[col]: sheet.cell_value(row_index, col) for col in range(min(sheet.ncols, len(headers)))}
            if not str(row.get("交易日期", "")).strip():
                continue
            amount = str(row.get("交易金额", "")).strip()
            if not parse_money_token(amount):
                continue
            transactions.append(
                make_transaction(
                    bank_key=str(manifest_item.get("bank_key", "ccb") or "ccb"),
                    bank_name="建设银行",
                    account_full_name=account_full_name,
                    account_tail=account_tail,
                    transaction_time=_format_yyyymmdd(str(row.get("交易日期", "")).strip()),
                    direction="inflow" if not amount.startswith("-") else "outflow",
                    amount=amount,
                    summary=str(row.get("摘要", "")).strip(),
                    balance=str(row.get("账户余额", "")).strip(),
                    counterparty=str(row.get("对方账号与户名", "")).strip(),
                    channel=str(row.get("交易地点/附言", "")).strip(),
                    source_records=[
                        {
                            "source_type": "email_attachment_xls",
                            "source_file": str(path),
                            "message_uid": manifest_item.get("message_uid", ""),
                            "message_id": manifest_item.get("message_id", ""),
                            "sheet": sheet.name,
                            "row": row_index + 1,
                        }
                    ],
                    confidence=0.92,
                    raw_record=row,
                )
            )
    return transactions


def _read_xlsx_transactions(path: Path, manifest_item: dict[str, Any]) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    transactions: list[dict[str, Any]] = []
    try:
        for sheet in book.worksheets:
            values = list(sheet.iter_rows(values_only=True))
            header_row = _find_sequence_header_row(values)
            if header_row is None:
                continue
            headers = [str(value or "").strip() for value in values[header_row]]
            rows = [
                {headers[index]: str(value or "").strip() for index, value in enumerate(row) if index < len(headers)}
                for row in values[header_row + 1 :]
            ]
            transactions.extend(
                _read_generic_bank_rows(
                    path,
                    manifest_item,
                    rows,
                    "email_attachment_xlsx",
                    sheet_name=sheet.title,
                    first_data_row=header_row + 2,
                )
            )
    finally:
        book.close()
    return transactions


def _read_generic_bank_rows(
    path: Path,
    manifest_item: dict[str, Any],
    rows: list[dict[str, str]],
    source_type: str,
    sheet_name: str = "",
    first_data_row: int = 2,
) -> list[dict[str, Any]]:
    transactions: list[dict[str, Any]] = []
    for offset, row in enumerate(rows):
        transaction_time = _first_value(row, "交易时间", "交易日期", "交易日", "发生时间", "记账时间")
        posting_date = _first_value(row, "入账日期", "记账日期", "入账日")
        debit = _clean_money(_first_value(row, "支出金额", "借方发生额", "借方金额"))
        credit = _clean_money(_first_value(row, "收入金额", "贷方发生额", "贷方金额"))
        amount = _clean_money(_first_value(row, "交易金额", "金额", "发生额"))
        direction = _direction_from_cn(_first_value(row, "收/支", "收支", "交易方向", "借贷标志"))
        if parse_money_token(debit) and debit not in {"0", "0.0", "0.00"}:
            amount, direction = debit, "outflow"
        elif parse_money_token(credit) and credit not in {"0", "0.0", "0.00"}:
            amount, direction = credit, "inflow"
        if not parse_money_token(amount) or _is_non_transaction_row(row, amount):
            continue
        if direction == "unknown":
            direction = "outflow" if amount.startswith("-") else "inflow" if amount.startswith("+") else "unknown"
        transaction_type = _first_value(row, "交易类型", "业务类型", "交易种类", "业务种类", "交易类别")
        summary = _first_value(row, "摘要", "交易摘要", "备注", "附言") or transaction_type
        combined_counterparty = _first_value(row, "对方账号与户名", "对方户名与账号")
        combined_name, combined_account = _split_counterparty(combined_counterparty)
        counterparty = _first_value(
            row,
            "交易对方",
            "对方户名",
            "对方名称",
            "收款方",
            "收款方户名",
            "付款方",
            "付款方户名",
        ) or combined_name
        counterparty_account = _first_value(
            row,
            "对方账号",
            "对方账户",
            "对方卡号",
            "收款账号",
            "收款方账号",
            "付款账号",
            "付款方账号",
        ) or combined_account
        if not counterparty and combined_counterparty:
            counterparty = combined_counterparty
        merchant = _first_value(row, "商户名称", "商户", "交易商户") or counterparty
        account_full_name = _first_value(row, "账户名称", "户名", "本方户名")
        account_number = _first_value(row, "卡号", "账号", "卡号/账号", "本方账号")
        account_display = account_full_name or _full_account_value(account_number)
        account_tail = _account_tail_from_text(account_number)
        source = {
            "source_type": source_type,
            "source_file": str(path),
            "message_uid": manifest_item.get("message_uid", ""),
            "message_id": manifest_item.get("message_id", ""),
            "row": first_data_row + offset,
        }
        if sheet_name:
            source["sheet"] = sheet_name
        transactions.append(
            make_transaction(
                bank_key=str(manifest_item.get("bank_key", "unknown") or "unknown"),
                bank_name=_bank_name(str(manifest_item.get("bank_key", ""))),
                account_full_name=account_display,
                account_tail=account_tail,
                transaction_time=_normalize_date_time(transaction_time),
                posting_date=_normalize_date_time(posting_date)[:10],
                direction=direction,
                amount=amount,
                currency=_first_value(row, "币种", "货币") or "CNY",
                merchant=merchant,
                counterparty=counterparty,
                counterparty_account=counterparty_account,
                summary=summary,
                balance=_clean_money(_first_value(row, "账户余额", "余额")),
                channel=_first_value(
                    row,
                    "交易渠道",
                    "渠道",
                    "交易渠道名称",
                    "交易场所",
                    "交易地点",
                    "交易地点/附言",
                    "交易网点",
                    "交易机构",
                ),
                transaction_type=transaction_type,
                transaction_reference=_first_value(row, "流水号", "交易流水号", "参考号", "交易序号"),
                source_records=[source],
                confidence=0.9,
                raw_record=row,
            )
        )
    return transactions


def _read_attachment_with_ai(
    path: Path,
    manifest_item: dict[str, Any],
    ai_fallback: FinancialDocumentAiFallback,
) -> list[dict[str, Any]]:
    text = _attachment_text(path)
    source = {
        "source_type": f"email_attachment_{path.suffix.lower().lstrip('.')}",
        "source_file": str(path),
        "message_uid": manifest_item.get("message_uid", ""),
        "message_id": manifest_item.get("message_id", ""),
    }
    bank_name = str(manifest_item.get("bank_name", "")).strip()
    bank_key = str(manifest_item.get("bank_key", "unknown") or "unknown")
    bank_key = {
        "交通银行": "bocom",
        "工商银行": "icbc",
        "建设银行": "ccb",
        "招商银行": "cmb",
    }.get(bank_name, bank_key)
    if path.suffix.lower() == ".pdf":
        return ai_fallback.parse_pdf(
            path=path,
            text=text,
            bank_key=bank_key,
            bank_name=bank_name or _bank_name(bank_key),
            source_record=source,
            source_label=source["source_type"],
        )
    return ai_fallback.parse(
        text=text,
        bank_key=bank_key,
        bank_name=bank_name or _bank_name(bank_key),
        source_record=source,
        source_label=source["source_type"],
    )


def _attachment_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(_fix_mojibake(page.extract_text() or "") for page in PdfReader(str(path)).pages)
    if suffix == ".csv":
        return _read_csv_text(path)
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        return "\n".join(
            "\t".join(str(sheet.cell_value(row, col)) for col in range(sheet.ncols))
            for sheet in book.sheets()
            for row in range(min(sheet.nrows, 500))
        )
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return "\n".join(
                "\t".join(str(value or "") for value in row)
                for sheet in book.worksheets
                for row in list(sheet.iter_rows(values_only=True))[:500]
            )
        finally:
            book.close()
    return ""


def _find_header_row(sheet: Any) -> int | None:
    for row_index in range(sheet.nrows):
        values = [str(sheet.cell_value(row_index, col)).strip() for col in range(sheet.ncols)]
        if "交易日期" in values and "交易金额" in values:
            return row_index
    return None


def _find_sequence_header_row(rows: list[tuple[Any, ...]]) -> int | None:
    date_aliases = {"交易时间", "交易日期", "交易日", "发生时间", "记账时间"}
    amount_aliases = {"交易金额", "金额", "发生额", "支出金额", "收入金额", "借方发生额", "贷方发生额"}
    for row_index, row in enumerate(rows[:30]):
        values = {str(value or "").strip() for value in row}
        if values & date_aliases and values & amount_aliases:
            return row_index
    return None


def _first_value(row: dict[str, str], *aliases: str) -> str:
    normalized = {re.sub(r"\s+", "", str(key)): str(value or "").strip() for key, value in row.items()}
    for alias in aliases:
        if alias in normalized and normalized[alias]:
            return normalized[alias]
    return ""


def _split_counterparty(value: str) -> tuple[str, str]:
    text = str(value or "").strip()
    if not text:
        return "", ""
    account_match = re.search(r"(?<!\d)(?:\d[\d*\s-]{5,}\d|\*{4,}\d{2,})(?!\d)", text)
    if not account_match:
        return text, ""
    account = re.sub(r"[\s-]", "", account_match.group(0))
    name = (text[: account_match.start()] + " " + text[account_match.end() :]).strip(" /,，;；")
    return re.sub(r"\s+", " ", name), account


def _is_non_transaction_row(row: dict[str, str], amount: str) -> bool:
    text = " ".join(str(value or "") for value in row.values())
    if any(term in text for term in ("合计", "总计", "本期应还", "最低还款", "信用额度", "可用额度")):
        return True
    try:
        return float(amount.replace(",", "")) == 0
    except ValueError:
        return True


def _normalize_date_time(value: str) -> str:
    text = str(value or "").strip()
    digits = re.sub(r"\D", "", text)
    if len(digits) == 8 and len(text) <= 10:
        return f"{digits[:4]}-{digits[4:6]}-{digits[6:8]}"
    match = re.match(
        r"(?P<year>20\d{2})[-/.年](?P<month>\d{1,2})[-/.月](?P<day>\d{1,2})日?"
        r"(?:\s+(?P<time>\d{1,2}:\d{2}(?::\d{2})?))?",
        text,
    )
    if not match:
        return text
    date = f"{int(match.group('year')):04d}-{int(match.group('month')):02d}-{int(match.group('day')):02d}"
    return f"{date} {match.group('time')}" if match.group("time") else date


def _bank_name(bank_key: str) -> str:
    return {
        "icbc": "工商银行",
        "cmb": "招商银行",
        "ccb": "建设银行",
        "bocom": "交通银行",
    }.get(bank_key, "")


def _account_full_from_xls(sheet: Any) -> str:
    for row_index in range(min(sheet.nrows, 5)):
        for col in range(sheet.ncols):
            value = str(sheet.cell_value(row_index, col))
            match = ACCOUNT_RE.search(value)
            if match:
                return _full_account_value(match.group(1))
    return ""


def _account_tail_from_lines(lines: list[str]) -> str:
    for line in lines[:20]:
        match = ACCOUNT_RE.search(line)
        if match:
            return match.group(1)[-4:]
    return ""


def _account_tail_from_text(value: str) -> str:
    match = CARD_TAIL_RE.search(value or "")
    if match:
        return match.group(1)[-4:]
    match = ACCOUNT_RE.search(value or "")
    if match:
        return match.group(1)[-4:]
    digits = re.sub(r"\D", "", value or "")
    if len(digits) >= 4:
        return digits[-4:]
    return ""


def _full_account_value(value: str) -> str:
    digits = re.sub(r"\D", "", value or "")
    return value if len(digits) > 4 and "*" not in value else ""


def _account_full_from_lines(lines: list[str]) -> str:
    for line in lines[:20]:
        match = ACCOUNT_RE.search(line)
        if match:
            value = match.group(1)
            return _full_account_value(value)
    return ""


def _direction_from_cn(value: str) -> str:
    text = (value or "").strip()
    if "收入" in text or text == "收":
        return "inflow"
    if "支出" in text or text == "支":
        return "outflow"
    return "unknown"


def _clean_money(value: str) -> str:
    return (
        str(value or "")
        .strip()
        .replace("￥", "")
        .replace("¥", "")
        .replace("\xa5", "")
        .replace("元", "")
        .replace(",", "")
    )


def _format_yyyymmdd(value: str) -> str:
    digits = re.sub(r"\D", "", value)
    if len(digits) == 8:
        return f"{digits[:4]}-{digits[4:6]}-{digits[6:8]}"
    return value


def _fix_mojibake(value: str) -> str:
    try:
        fixed = value.encode("latin1").decode("gbk")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return value
    return fixed if _score_cjk(fixed) > _score_cjk(value) else value


def _score_cjk(value: str) -> int:
    return sum(1 for char in value if "\u4e00" <= char <= "\u9fff")
