from __future__ import annotations

import logging
import csv
import re
import tempfile
from io import StringIO
from pathlib import Path
from typing import Any, Mapping

from flows.modules.json_extractor import parse_json_from_text
from flows.modules.llamacpp_client import LlamaCppClient, LlamaCppConfig
from flows.modules.order_image_prompt import build_order_image_prompt
from flows.modules.order_schema import make_order


logger = logging.getLogger(__name__)

ORDER_HEADER_RE = re.compile(
    r"(?P<time>20\d{2}-\d{2}-\d{2}\s+\d{2}:\d{2}:\d{2})\s*"
    r"订单号[：:]\s*(?P<order_id>\d{6,})"
)
LOOSE_ORDER_DATE_RE = re.compile(r"20\d{2}[-/]\d{1,2}[-/]\d{1,2}\s+\d{1,2}:\d{2}(?::\d{2})?")
MONEY_RE = re.compile(r"[¥￥]\s*([\d,]+(?:\.\d{1,2})?)")
QUANTITY_RE = re.compile(r"\bx\s*(\d+)\b", re.IGNORECASE)
STATUS_VALUES = (
    "已完成",
    "已取消",
    "已拆分",
    "待付款",
    "待收货",
    "待评价",
    "订单关闭",
    "等待付款",
    "正在出库",
)
JD_NAVIGATION_LABELS = {
    "资产中心",
    "客户服务",
    "企业租赁",
    "海外房产预约",
    "流量加油站",
    "我的京东国际",
    "医药服务",
    "小金库",
    "京东白条",
    "京东机票",
    "交易纠纷",
    "个人信息",
    "收货地址",
}
JD_NAVIGATION_PREFIX_RE = re.compile(
    r"^(?:优惠券\s*(?:\(\d+\))?|礼品卡\s*(?:\(\d+\))?|我的发票|返修退换货|"
    r"价格保护|我的问答|购买咨询|京东维修|举报中心|设置|关注的商品|"
    r"关注的店铺|关注的活动|京豆|银行卡|红包|领货码|定期购|京东通信)\s*"
)


def read_jd_pdf_orders(
    raw_root: Path,
    config: Mapping[str, Any] | None = None,
    project_root: Path | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """读取 raw_data/jd 中的京东订单 PDF。

    先使用 PDF 自带文本层；仅在字体编码导致无法稳定读取时，对该页调用
    本地视觉模型。
    """
    from pypdf import PdfReader

    directory = raw_root / "jd"
    orders: list[dict[str, Any]] = []
    failures: list[str] = []
    files_seen = 0
    pages_seen = 0
    ai_pages = 0
    ai = _LazyOrderVision(config, project_root) if config is not None and project_root is not None else None

    for path in sorted(directory.glob("*.pdf")):
        files_seen += 1
        try:
            reader = PdfReader(str(path))
            for page_number, page in enumerate(reader.pages, start=1):
                pages_seen += 1
                text = page.extract_text(extraction_mode="layout") or ""
                parsed = parse_jd_pdf_page(path, page_number, text)
                if not parsed and ai is not None and _requires_visual_fallback(text):
                    parsed = ai.parse_pdf_page(path, page_number)
                    ai_pages += 1
                orders.extend(parsed)
        except Exception as exc:
            failures.append(f"{path}: {type(exc).__name__}: {exc}")
            logger.warning("Failed parsing JD order PDF: file=%s error=%s", path, exc)

    return orders, {
        "directory": str(directory),
        "files_seen": files_seen,
        "pages_seen": pages_seen,
        "orders_seen": len(orders),
        "ai_pages": ai_pages,
        "failures": failures,
    }


def read_order_spreadsheets(
    raw_root: Path,
    platforms: list[str],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    orders: list[dict[str, Any]] = []
    failures: list[str] = []
    files_seen = 0
    for platform in platforms:
        directory = raw_root / platform
        paths = sorted(
            path
            for path in directory.rglob("*")
            if path.is_file() and path.suffix.lower() in {".csv", ".xls", ".xlsx"}
        )
        for path in paths:
            files_seen += 1
            try:
                for rows, sheet_name, first_data_row in _spreadsheet_rows(path):
                    orders.extend(
                        parse_order_rows(path, platform, rows, sheet_name, first_data_row)
                    )
            except Exception as exc:
                failures.append(f"{path}: {type(exc).__name__}: {exc}")
                logger.warning("Failed parsing order spreadsheet: file=%s error=%s", path, exc)
    return orders, {"files_seen": files_seen, "orders_seen": len(orders), "failures": failures}


def parse_order_rows(
    path: Path,
    platform: str,
    rows: list[dict[str, str]],
    sheet_name: str = "",
    first_data_row: int = 2,
) -> list[dict[str, Any]]:
    orders: list[dict[str, Any]] = []
    for offset, row in enumerate(rows):
        title = _first_value(row, "商品全名", "商品名称", "商品标题", "订单标题", "商品")
        order_id = _first_value(row, "订单号", "订单编号", "交易号", "交易订单号")
        if not title and not order_id:
            continue
        source = {
            "source_type": "order_spreadsheet",
            "source_file": str(path),
            "row": first_data_row + offset,
            "order_index": offset,
        }
        if sheet_name:
            source["sheet"] = sheet_name
        raw_platform = _first_value(row, "平台", "来源平台") or platform
        orders.append(
            make_order(
                platform=raw_platform,
                source_file=str(path),
                order_index=offset,
                merchant=_first_value(row, "商户", "商家", "店铺", "商家名称", "店铺名称"),
                status=_first_value(row, "订单状态", "交易状态", "状态"),
                title=title,
                spec=_first_value(row, "商品规格", "规格", "型号"),
                quantity=_first_value(row, "购买数量", "数量"),
                paid_amount=_first_value(row, "实付金额", "实付款", "支付金额", "订单金额", "金额"),
                original_amount=_first_value(row, "原价", "商品金额", "应付金额"),
                shipping_fee=_first_value(row, "运费", "配送费"),
                order_time=_first_value(row, "下单时间", "订单时间", "交易时间", "创建时间"),
                order_id=order_id,
                logistics=_first_value(row, "物流", "物流信息"),
                confidence=0.96,
                raw_record=row,
                source_type="order_spreadsheet",
                source_records=[source],
            )
        )
    return orders


def parse_jd_pdf_page(path: Path, page_number: int, text: str) -> list[dict[str, Any]]:
    matches = list(ORDER_HEADER_RE.finditer(text))
    orders: list[dict[str, Any]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        block = text[match.start() : end]
        money = MONEY_RE.search(block)
        if not money:
            continue
        amount = money.group(1).replace(",", "")
        header_line = block.splitlines()[0]
        header_tail = header_line[match.end() - match.start() :].strip()
        merchant = "京东" if "京东" in header_tail or "您订单中的商品" in header_tail else header_tail
        title, quantity = _jd_title_and_quantity(block)
        status = next((value for value in STATUS_VALUES if value in block), "")
        orders.append(
            make_order(
                platform="jd",
                source_file=str(path),
                order_index=index,
                merchant=merchant or "京东",
                status=status,
                title=title,
                quantity=quantity,
                paid_amount=amount,
                order_time=match.group("time"),
                order_id=match.group("order_id"),
                confidence=0.96,
                warnings=[] if title else ["missing_title"],
                raw_record={"page_number": page_number, "text": _compact_text(block)},
                source_type="order_pdf",
                is_container=title == "拆分订单",
                source_records=[
                    {
                        "source_type": "order_pdf",
                        "source_file": str(path),
                        "page_number": page_number,
                        "order_index": index,
                    }
                ],
            )
        )
    return orders


def _requires_visual_fallback(text: str) -> bool:
    """只把确实无法读取的订单页交给视觉模型。

    京东导出的 PDF 常在最后附带只有页脚或订单续行的页面；这些页面可能含有
    乱码文本，但没有独立订单头，不能仅因规则解析结果为空就调用 AI。
    """
    if not text.strip():
        return True
    return bool(LOOSE_ORDER_DATE_RE.search(text))


def _jd_title_and_quantity(block: str) -> tuple[str, str]:
    for line in block.splitlines()[1:]:
        if not MONEY_RE.search(line):
            continue
        before_money = MONEY_RE.split(line, maxsplit=1)[0]
        quantity = QUANTITY_RE.search(before_money)
        if not quantity:
            continue
        raw_title = before_money[: quantity.start()].strip()
        raw_title = JD_NAVIGATION_PREFIX_RE.sub("", raw_title)
        parts = [part.strip() for part in re.split(r"\s{2,}", raw_title) if part.strip()]
        # 页面左侧导航栏与商品处在同一文本行，丢弃最左的导航单元。
        title_parts = parts[1:] if parts and parts[0] in JD_NAVIGATION_LABELS else parts
        title = " ".join(title_parts)
        title = re.sub(r"\s+", " ", title).strip()
        title = re.sub(r"^(?:收货人[：:].*?\s+)?", "", title)
        return title, quantity.group(1)
    if "订单金额" in block and "已拆分" in block:
        return "拆分订单", ""
    return "", ""


def _compact_text(value: str) -> str:
    return "\n".join(re.sub(r"\s+", " ", line).strip() for line in value.splitlines() if line.strip())


def _spreadsheet_rows(path: Path) -> list[tuple[list[dict[str, str]], str, int]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        data = path.read_bytes()
        text = next(
            (data.decode(encoding) for encoding in ("utf-8-sig", "gb18030", "gbk") if _can_decode(data, encoding)),
            data.decode("latin1"),
        )
        lines = [line for line in text.splitlines() if line.strip()]
        header_index = _header_index([next(csv.reader([line])) for line in lines])
        if header_index is None:
            return []
        reader = csv.DictReader(StringIO("\n".join(lines[header_index:])))
        rows = [{str(key or "").strip(): str(value or "").strip() for key, value in row.items()} for row in reader]
        return [(rows, "", header_index + 2)]
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        result = []
        for sheet in book.sheets():
            values = [[sheet.cell_value(row, col) for col in range(sheet.ncols)] for row in range(sheet.nrows)]
            result.extend(_sequence_rows(values, sheet.name))
        return result
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        result = []
        for sheet in book.worksheets:
            result.extend(_sequence_rows(list(sheet.iter_rows(values_only=True)), sheet.title))
        return result
    finally:
        book.close()


def _sequence_rows(values: list[list[Any] | tuple[Any, ...]], sheet_name: str) -> list[tuple[list[dict[str, str]], str, int]]:
    header_index = _header_index(values)
    if header_index is None:
        return []
    headers = [str(value or "").strip() for value in values[header_index]]
    rows = [
        {headers[index]: str(value or "").strip() for index, value in enumerate(row) if index < len(headers)}
        for row in values[header_index + 1 :]
    ]
    return [(rows, sheet_name, header_index + 2)]


def _header_index(values: list[list[Any] | tuple[Any, ...]]) -> int | None:
    aliases = {"订单号", "订单编号", "商品名称", "商品全名", "下单时间", "实付金额"}
    for index, row in enumerate(values[:30]):
        if {str(value or "").strip() for value in row} & aliases:
            return index
    return None


def _first_value(row: dict[str, str], *aliases: str) -> str:
    normalized = {re.sub(r"\s+", "", str(key)): str(value or "").strip() for key, value in row.items()}
    return next((normalized[alias] for alias in aliases if normalized.get(alias)), "")


def _can_decode(data: bytes, encoding: str) -> bool:
    try:
        data.decode(encoding)
        return True
    except UnicodeDecodeError:
        return False


class _LazyOrderVision:
    def __init__(self, config: Mapping[str, Any], project_root: Path) -> None:
        self.client = LlamaCppClient(LlamaCppConfig.from_config(dict(config), project_root))
        self.checked = False

    def parse_pdf_page(self, path: Path, page_number: int) -> list[dict[str, Any]]:
        if not self.checked:
            _health, models = self.client.ensure_server()
            self.client.assert_model_available(models)
            self.checked = True
        with tempfile.TemporaryDirectory(prefix="jd_order_ocr_") as directory:
            image = _render_pdf_page(path, page_number, Path(directory))
            response = self.client.chat_with_image(
                build_order_image_prompt("jd", f"{path.name}#page={page_number}"),
                image,
                max_tokens=4096,
            )
        payload = parse_json_from_text(response)
        rows = payload.get("orders", []) if isinstance(payload, dict) else []
        orders: list[dict[str, Any]] = []
        for index, row in enumerate(rows if isinstance(rows, list) else []):
            if not isinstance(row, dict):
                continue
            orders.append(
                make_order(
                    platform="jd",
                    source_file=str(path),
                    order_index=index,
                    merchant=row.get("merchant", ""),
                    status=row.get("status", ""),
                    title=row.get("title", ""),
                    spec=row.get("spec", ""),
                    quantity=row.get("quantity", ""),
                    paid_amount=row.get("paid_amount", ""),
                    original_amount=row.get("original_amount", ""),
                    shipping_fee=row.get("shipping_fee", ""),
                    order_time=row.get("order_time", ""),
                    order_id=row.get("order_id", ""),
                    logistics=row.get("logistics", ""),
                    actions=row.get("actions", []) if isinstance(row.get("actions"), list) else [],
                    is_partial=bool(row.get("is_partial", False)),
                    confidence=row.get("confidence", 0.68),
                    warnings=["parsed_by_local_ai", "requires_human_review"],
                    notes=row.get("notes", ""),
                    raw_record=row,
                    source_type="order_pdf_ai",
                    source_records=[
                        {
                            "source_type": "order_pdf_ai",
                            "source_file": str(path),
                            "page_number": page_number,
                            "order_index": index,
                        }
                    ],
                )
            )
        logger.info("Local AI parsed JD PDF page: file=%s page=%s orders=%s", path, page_number, len(orders))
        return orders


def _render_pdf_page(path: Path, page_number: int, output_dir: Path) -> Path:
    import fitz

    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"page_{page_number:04d}.png"
    with fitz.open(path) as document:
        page = document[page_number - 1]
        page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(output)
    return output
